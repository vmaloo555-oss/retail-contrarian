"""
strategy_pdf.py -- render the option-strategy sheet as a clean one-look PDF.

Only actionable rows are tabled (a direction AND real legs). Flat / thin-chain /
no-chain names are summarised as counts at the end rather than padding 200 rows.

    python strategy_pdf.py "output/netpos 8 sep-n.xlsx"
    -> "output/strategy 8 sep-n.pdf"
"""
from __future__ import annotations
import os, re, sys, datetime as dt

import openpyxl
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph,
                                Spacer, Table, TableStyle)

INK      = colors.HexColor("#14181F")
MUTED    = colors.HexColor("#6B7280")
HAIRLINE = colors.HexColor("#E3E6EA")
BAND     = colors.HexColor("#F5F6F8")
BULL     = colors.HexColor("#1B7F5A")
BULL_BG  = colors.HexColor("#E8F4EF")
BEAR     = colors.HexColor("#B3341F")
BEAR_BG  = colors.HexColor("#FBEDEA")
PAPER    = colors.HexColor("#FFFFFF")

BODY  = ParagraphStyle("body", fontName="Helvetica", fontSize=7.6, leading=9.6,
                       textColor=INK, alignment=TA_LEFT)
LEGS  = ParagraphStyle("legs", parent=BODY, fontName="Helvetica-Bold")
SMALL = ParagraphStyle("small", parent=BODY, fontSize=6.9, leading=8.6,
                       textColor=MUTED)
H1    = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=16, leading=19,
                       textColor=INK)
SUB   = ParagraphStyle("sub", fontName="Helvetica", fontSize=8.4, leading=11,
                       textColor=MUTED)
SECT  = ParagraphStyle("sect", fontName="Helvetica-Bold", fontSize=9.6,
                       leading=12, textColor=colors.white)


def esc(x):
    """Escape for reportlab's mini-HTML — symbols like GVT&D, M&M, BAJAJ-AUTO."""
    return (str(x).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if x not in (None, "") else "")


def _fmt(x, nd=2):
    if x is None or x == "":
        return ""
    if isinstance(x, (int, float)):
        return f"{x:,.{nd}f}".rstrip("0").rstrip(".") if nd else f"{x:,.0f}"
    return esc(x)


def _read(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["option strategy"]
    hdr = [c.value for c in ws[1]]
    ix = {h: i for i, h in enumerate(hdr)}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = {h: r[i] for h, i in ix.items()}
        # the sheet header is "setup (retail strike map)" - normalise it
        for h in list(d):
            if isinstance(h, str) and h.startswith("setup"):
                d["setup"] = d[h]
        rows.append(d)
    wb.close()
    return rows


def _reading_name(xlsx):
    m = re.search(r"netpos (.+)\.xlsx$", os.path.basename(xlsx))
    return m.group(1) if m else os.path.basename(xlsx)


def build(xlsx, out_pdf=None, preview_txt=None):
    rows = _read(xlsx)
    name = _reading_name(xlsx)
    out_pdf = out_pdf or os.path.join(os.path.dirname(xlsx), f"strategy {name}.pdf")

    def actionable(r):
        return bool(r.get("legs")) and r.get("strategy") not in (None, "", "none")

    GROUPS = [
        ("BEARISH", "Retail's option book is net BULLISH here - long calls that "
                    "decay and short puts that get run over. Fade it short."),
        ("BULLISH", "Retail's option book is net BEARISH here - short calls that "
                    "get run over and long puts that decay. Fade it long."),
    ]
    groups = {k: [] for k, _ in GROUPS}
    for r in rows:
        if actionable(r) and (r.get("verdict") or "") in groups:
            groups[r["verdict"]].append(r)
    for k in groups:
        groups[k].sort(key=lambda r: -(float(r.get("conviction") or 0)))
    thin = sum(1 for r in rows if "thin" in (r.get("contrarian read") or ""))
    nochain = sum(1 for r in rows if "no option chain" in (r.get("contrarian read") or ""))
    flat = sum(1 for r in rows if (r.get("verdict") or "") == "balanced")
    nolegs = sum(1 for r in rows if not actionable(r)
                 and (r.get("verdict") or "") in ("BULLISH", "BEARISH"))

    regime = ""
    if preview_txt and os.path.exists(preview_txt):
        for ln in open(preview_txt, encoding="utf-8"):
            if ln.startswith("# NIFTY"):
                regime = ln.lstrip("# ").strip()
                break

    W, H = landscape(A4)
    doc = BaseDocTemplate(out_pdf, pagesize=(W, H),
                          leftMargin=13 * mm, rightMargin=13 * mm,
                          topMargin=12 * mm, bottomMargin=13 * mm,
                          title=f"Contrarian option strategies — {name}",
                          author="retail-contrarian")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")

    def furniture(canv, d):
        canv.saveState()
        canv.setFillColor(PAPER)
        canv.rect(0, 0, W, H, fill=1, stroke=0)
        canv.setStrokeColor(HAIRLINE); canv.setLineWidth(0.5)
        canv.line(d.leftMargin, 9 * mm, W - d.rightMargin, 9 * mm)
        canv.setFont("Helvetica", 6.8); canv.setFillColor(MUTED)
        canv.drawString(d.leftMargin, 6 * mm,
                        "Mechanical output from retail positioning — not advice. "
                        "Accumulation phase: NOT for trading. You place every order.")
        canv.drawRightString(W - d.rightMargin, 6 * mm, f"page {canv.getPageNumber()}")
        canv.restoreState()

    doc.addPageTemplates([PageTemplate(id="p", frames=[frame], onPage=furniture)])

    COLS = [24 * mm, 16 * mm, 14 * mm, 44 * mm, 46 * mm, 46 * mm, 38 * mm, 40 * mm]
    HEAD = ["Symbol", "Spot", "Conv", "Retail's option book (weighted)",
            "What it means", "Legs", "Alternative", "Caution"]

    def section(title, data, accent, accent_bg, blurb):
        if not data:
            return []
        bar = Table([[Paragraph(f"{title}  ·  {len(data)} names", SECT)]],
                    colWidths=[sum(COLS)], rowHeights=[7.6 * mm])
        bar.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), accent),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        body = [[Paragraph(f"<b>{h}</b>", SMALL) for h in HEAD]]
        for r in data:
            body.append([
                Paragraph(f"<b>{esc(r['symbol'])}</b>", BODY),
                Paragraph(_fmt(r.get("spot")), BODY),
                Paragraph(_fmt(r.get("conviction"), 2), BODY),
                Paragraph(esc(r.get("activity breakdown")), SMALL),
                Paragraph(esc(r.get("contrarian read")), SMALL),
                Paragraph(esc(r.get("legs")), LEGS),
                Paragraph(esc(r.get("alternative")), SMALL),
                Paragraph(f"<font color='#B3341F'>{esc(r.get('caution'))}</font>"
                          if r.get("caution") else "", SMALL)])
        t = Table(body, colWidths=COLS, repeatRows=1)
        st = [("VALIGN", (0, 0), (-1, -1), "TOP"),
              ("TOPPADDING", (0, 0), (-1, -1), 3.2),
              ("BOTTOMPADDING", (0, 0), (-1, -1), 3.2),
              ("LEFTPADDING", (0, 0), (-1, -1), 5),
              ("RIGHTPADDING", (0, 0), (-1, -1), 5),
              ("BACKGROUND", (0, 0), (-1, 0), BAND),
              ("LINEBELOW", (0, 0), (-1, 0), 0.6, HAIRLINE),
              ("LINEBELOW", (0, 1), (-1, -2), 0.25, HAIRLINE),
              ("TEXTCOLOR", (2, 1), (2, -1), accent)]
        for i in range(1, len(body)):
            if i % 2 == 0:
                st.append(("BACKGROUND", (0, i), (-1, i), accent_bg))
        t.setStyle(TableStyle(st))
        return [Paragraph(blurb, SMALL), Spacer(1, 2.2 * mm), bar, Spacer(1, 1.4 * mm),
                t, Spacer(1, 6 * mm)]

    story = [Paragraph("Contrarian option strategies", H1),
             Paragraph(f"Reading <b>{name}</b> &nbsp;·&nbsp; generated "
                       f"{dt.datetime.now():%d %b %Y %H:%M} IST"
                       + (f" &nbsp;·&nbsp; {regime}" if regime else ""), SUB),
             Spacer(1, 5 * mm)]
    PALETTE = {"BEARISH": (BEAR, BEAR_BG), "BULLISH": (BULL, BULL_BG)}
    actionable_n = 0
    for key, blurb in GROUPS:
        data = groups[key]
        actionable_n += len(data)
        acc, bg = PALETTE[key]
        story += section(key, data, acc, bg, blurb)

    tot = len(rows)
    excl = Table([[Paragraph(
        f"<b>Not shown</b> &nbsp; {thin} thin option chain &nbsp;·&nbsp; {flat} balanced book, "
        f"two-sided (no edge) &nbsp;·&nbsp; {nolegs} direction but no usable strikes "
        f"&nbsp;·&nbsp; {nochain} no chain &nbsp; — of {tot} symbols, "
        f"<b>{actionable_n} actionable</b>.", SMALL)]],
        colWidths=[sum(COLS)])
    excl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), BAND),
                              ("LEFTPADDING", (0, 0), (-1, -1), 6),
                              ("TOPPADDING", (0, 0), (-1, -1), 5),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story.append(excl)
    doc.build(story)
    return out_pdf


if __name__ == "__main__":
    xlsx = sys.argv[1]
    prev = sys.argv[2] if len(sys.argv) > 2 else \
        os.path.join(os.path.dirname(xlsx), f"preview {_reading_name(xlsx)}.txt")
    print(build(xlsx, preview_txt=prev))
