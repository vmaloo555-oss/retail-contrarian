"""
pos_parser.py -- Parser for the twice-daily NSE F&O retail-positioning feed.

The feed is an .xlsx workbook with a sheet named `pos`, one row per F&O symbol.
Two files arrive daily (~11:30 and ~14:30 IST); the 14:30 file drives decisions.
This parser does not care which file it is handed -- it parses whatever `pos`
sheet is in the path you give it.

0-indexed column map (per spec):
    0  symbol
    1  sector
    2  LTP (last traded price)
    3  retail FUT  long  current     4  retail FUT  short current
    5  retail FUT  long  previous    6  retail FUT  short previous
    7  retail CALL long  current     8  retail CALL short current
    9  retail CALL long  previous   10  retail CALL short previous
   11  retail PUT  buy   current    12  retail PUT  sell  current
   13  retail PUT  buy   previous   14  retail PUT  sell  previous
   15  retail NET POSITION (current)  ==  (3 - 4) + (7 - 8) + (12 - 11)
   16  net position (previous)
   17  DIFF                          <- ignored
   18..N   net-position history, newest -> oldest (~2 readings / trading day).
           N is the last column whose date header is a date, capped at 34.
           Some exports append deltas / a repeated symbol / a price / a %
           change straight after the history; those are NOT history.
   >N      ignored

Public contract
---------------
parse_pos(path, sheet="pos", header_rows="auto", tol=1e-6) -> PosFeed

PosFeed.records : list[SymbolRecord], one per symbol row, in sheet order
PosFeed.mismatches : list[Mismatch], every row where col 15 != the formula
PosFeed.warnings : list[str], non-fatal data-quality notes

SymbolRecord fields:
    symbol      : str
    ltp         : float | None
    net_series  : list[float]   # [col 15] + the dated history cols, newest first
  (plus diagnostics: row, sector, net_current, net_current_expected, verified)

CLI
---
    python pos_parser.py FEED.xlsx              # human summary + mismatch report
    python pos_parser.py FEED.xlsx --json       # machine-readable dump
    python pos_parser.py FEED.xlsx --tol 0.5    # looser formula tolerance
Exit code is 0 when every row verifies, 1 when any mismatch is found.

Requires: openpyxl  (pip install openpyxl)
"""

from __future__ import annotations

import argparse
import json
import math
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterator, Optional

SHEET_DEFAULT = "pos"

# --- column indices (0-based) ------------------------------------------------
COL_SYMBOL = 0
COL_SECTOR = 1
COL_LTP = 2
COL_FUT_L_CUR, COL_FUT_S_CUR = 3, 4
COL_CALL_L_CUR, COL_CALL_S_CUR = 7, 8
COL_PUT_B_CUR, COL_PUT_S_CUR = 11, 12
COL_NET_CUR = 15
COL_NET_PREV = 16
COL_DIFF = 17  # ignored
COL_HIST_START, COL_HIST_END = 18, 34  # inclusive
MIN_COLS = COL_HIST_END + 1  # need columns 0..34 -> 35 cells

# inputs to the col-15 formula, in the order used by _expected_net()
FORMULA_INPUTS = (
    COL_FUT_L_CUR, COL_FUT_S_CUR,
    COL_CALL_L_CUR, COL_CALL_S_CUR,
    COL_PUT_B_CUR, COL_PUT_S_CUR,
)

_BLANK_TOKENS = {"", "-", "--", "–", "—", "na", "n/a", "#n/a",
                 "nil", "null", "none", "#value!", "#div/0!", "#ref!"}


def _num(v):
    """Coerce a cell value to float.

    Returns (value, ok):
      value -- float, or None for a blank / recognised null token
      ok    -- False only when the cell held non-blank text that is not a number
               (caller decides whether that is a hard error or just a warning)
    Handles thousands separators, unicode minus, and (parenthesised) negatives.
    """
    if v is None:
        return None, True
    if isinstance(v, bool):
        return float(v), True
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isnan(v):
            return None, True
        return float(v), True
    s = str(v).strip()
    if s.lower() in _BLANK_TOKENS:
        return None, True
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1].strip()
    s = s.replace(",", "").replace("−", "-").replace("–", "-").replace(" ", "")
    try:
        f = float(s)
    except ValueError:
        return None, False
    return (-f if neg else f), True


def _expected_net(vals):
    """vals: 6 floats in FORMULA_INPUTS order. Returns (3-4)+(7-8)+(12-11)."""
    fut_l, fut_s, call_l, call_s, put_b, put_s = vals
    return (fut_l - fut_s) + (call_l - call_s) + (put_s - put_b)


# --- result types ----------------------------------------------------------
@dataclass
class Mismatch:
    row: int                       # 1-based sheet row
    symbol: str
    expected: Optional[float]      # from the formula
    actual: Optional[float]        # value found in col 15 (None if blank/non-numeric)
    delta: Optional[float]         # actual - expected
    note: str = ""                 # e.g. blank/non-numeric inputs


@dataclass
class SymbolRecord:
    row: int
    symbol: str
    sector: Optional[str]
    ltp: Optional[float]
    net_series: list               # [col 15] + dated history cols, newest first
    net_current: Optional[float]   # raw col 15
    net_current_expected: Optional[float]
    verified: bool                 # col 15 matched the formula within tol

    def as_contract(self) -> dict:
        """The three fields the spec asks for."""
        return {"symbol": self.symbol, "ltp": self.ltp, "net_series": self.net_series}


@dataclass
class PosFeed:
    path: str
    sheet: str
    header_rows: int
    records: list = field(default_factory=list)
    mismatches: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.mismatches

    def __iter__(self) -> Iterator[SymbolRecord]:
        return iter(self.records)

    def __len__(self) -> int:
        return len(self.records)

    def find(self, symbol: str) -> Optional[SymbolRecord]:
        key = symbol.strip().upper()
        for r in self.records:
            if r.symbol.upper() == key:
                return r
        return None


# --- parser --------------------------------------------------------------
_HEADER_SYMBOL_TOKENS = {"", "symbol", "symbols", "ticker", "scrip", "scripname",
                         "name", "stock", "underlying", "instrument"}


def _is_text_nonnumber(v) -> bool:
    """True only for a non-blank cell that is text and not parseable as a number."""
    if v is None:
        return False
    if str(v).strip().lower() in _BLANK_TOKENS:
        return False
    _, ok = _num(v)
    return not ok


def _looks_like_header(r) -> bool:
    """Header row = a known symbol-column label, or several numeric columns
    carrying text labels at once. A lone corrupt cell in a data row does NOT
    trip this (that would silently drop a symbol)."""
    c0 = str(r[COL_SYMBOL]).strip().lower() if r[COL_SYMBOL] is not None else ""
    if c0 in _HEADER_SYMBOL_TOKENS:
        return True
    text_hits = sum(_is_text_nonnumber(r[c])
                    for c in (COL_LTP, COL_NET_CUR, COL_NET_PREV, COL_HIST_START))
    return text_hits >= 2


def _detect_header_rows(rows) -> int:
    """Count leading non-data rows (see _looks_like_header)."""
    n = 0
    for r in rows[:5]:
        rr = (list(r) if r else []) + [None] * MIN_COLS
        if _looks_like_header(rr):
            n += 1
            continue
        break
    return n


_DATE_TEXT = re.compile(r"^\s*\d{1,2}[/-]\d{1,2}([/-]\d{2,4})?\s*$")


def _is_date_header(v) -> bool:
    import datetime as _dt
    if isinstance(v, (_dt.datetime, _dt.date)):
        return True
    return isinstance(v, str) and bool(_DATE_TEXT.match(v))


def _hist_end(rows, hr: int) -> int:
    """Last real history column.

    The history block starts at COL_HIST_START and runs while the date header
    row still carries a date. Some exports append unrelated columns straight
    after it (deltas, a repeated symbol, a price, a % change); walking blindly
    to COL_HIST_END swallows those and a fractional % change then becomes the
    oldest 'position', which blows up every trend score. Stop at the first
    non-date header instead. Falls back to COL_HIST_END if no header is found,
    and never reads past it.
    """
    hdr = None
    for i in range(min(hr, len(rows))):
        r = rows[i] or ()
        if len(r) > COL_HIST_START and _is_date_header(r[COL_HIST_START]):
            hdr = r
            break
    if hdr is None:
        return COL_HIST_END
    end = COL_HIST_START
    for c in range(COL_HIST_START, min(len(hdr), COL_HIST_END + 1)):
        if not _is_date_header(hdr[c]):
            break
        end = c
    return end


def parse_pos(path, sheet: str = SHEET_DEFAULT, header_rows="auto",
              tol: float = 1e-6) -> PosFeed:
    """Load the `pos` sheet, verify col 15 on every row, return a PosFeed.

    header_rows : "auto" (default) sniffs leading non-data rows, or pass an int.
    tol         : max |actual - expected| for col 15 to count as verified.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise SystemExit("openpyxl is required:  pip install openpyxl") from e

    p = Path(path)
    if not p.exists():
        raise SystemExit(f"file not found: {p}")

    wb = load_workbook(p, data_only=True, read_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise SystemExit(
                f"sheet {sheet!r} not found; sheets present: {wb.sheetnames}")
        ws = wb[sheet]
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    hr = _detect_header_rows(rows) if header_rows == "auto" else int(header_rows)
    hist_end = _hist_end(rows, hr)
    feed = PosFeed(path=str(p), sheet=sheet, header_rows=hr)
    if hist_end < COL_HIST_END:
        feed.warnings.append(
            f"history block ends at column {hist_end} (not {COL_HIST_END}); "
            f"columns {hist_end + 1}..{COL_HIST_END} are not dated and were ignored")

    seen = {}
    any_net_present = False

    for idx in range(hr, len(rows)):
        raw = rows[idx] or ()
        sheet_row = idx + 1
        r = list(raw) + [None] * (MIN_COLS - len(raw))

        sym = r[COL_SYMBOL]
        sym = "" if sym is None else str(sym).strip()
        if sym == "":
            continue  # trailing / spacer row

        if sym.upper() in seen:
            feed.warnings.append(
                f"row {sheet_row}: duplicate symbol {sym!r} "
                f"(first seen row {seen[sym.upper()]})")
        else:
            seen[sym.upper()] = sheet_row

        sector = r[COL_SECTOR]
        sector = None if sector is None else str(sector).strip() or None

        ltp, ltp_ok = _num(r[COL_LTP])
        if not ltp_ok:
            feed.warnings.append(
                f"row {sheet_row} ({sym}): non-numeric LTP {r[COL_LTP]!r}")

        # --- verify col 15 against the formula ---
        parts, blank_inputs, bad_inputs = [], [], []
        for c in FORMULA_INPUTS:
            val, ok = _num(r[c])
            if not ok:
                bad_inputs.append(c)
            if val is None:
                blank_inputs.append(c)
                val = 0.0
            parts.append(val)
        expected = _expected_net(parts)

        actual, actual_ok = _num(r[COL_NET_CUR])
        if actual is not None:
            any_net_present = True

        verified = True
        if not actual_ok:
            verified = False
            feed.mismatches.append(Mismatch(
                sheet_row, sym, expected, None, None,
                f"col 15 non-numeric: {r[COL_NET_CUR]!r}"))
        elif actual is None:
            verified = False
            feed.mismatches.append(Mismatch(
                sheet_row, sym, expected, None, None, "col 15 blank"))
        else:
            delta = actual - expected
            if abs(delta) > tol:
                verified = False
                if bad_inputs:
                    note = f"non-numeric formula inputs at cols {bad_inputs}"
                elif blank_inputs:
                    note = f"blank formula inputs at cols {blank_inputs} treated as 0"
                else:
                    note = ""
                feed.mismatches.append(Mismatch(
                    sheet_row, sym, expected, actual, delta, note))

        # --- net-position series: [col 15] + cols 18..34, blanks dropped, order kept ---
        series = []
        for c in (COL_NET_CUR, *range(COL_HIST_START, hist_end + 1)):
            val, _ = _num(r[c])
            if val is not None:
                series.append(val)

        feed.records.append(SymbolRecord(
            row=sheet_row, symbol=sym, sector=sector, ltp=ltp,
            net_series=series, net_current=actual,
            net_current_expected=expected, verified=verified))

    if feed.records and not any_net_present:
        feed.warnings.append(
            "every col 15 is blank -- if the workbook stores this as a formula, "
            "openpyxl only sees cached results: open & save it in Excel, or have "
            "the feed export values.")

    return feed


# --- CLI -----------------------------------------------------------------
def _fmt(x, nd=2):
    return "-" if x is None else f"{x:,.{nd}f}"


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="Parse and verify the NSE F&O retail-positioning `pos` feed.")
    ap.add_argument("path", help="path to the feed .xlsx")
    ap.add_argument("--sheet", default=SHEET_DEFAULT)
    ap.add_argument("--header-rows", default="auto",
                    help='"auto" (default) or an integer count of leading rows')
    ap.add_argument("--tol", type=float, default=1e-6,
                    help="max |actual-expected| for col 15 to verify (default 1e-6)")
    ap.add_argument("--json", action="store_true",
                    help="emit records + mismatches + warnings as JSON")
    ap.add_argument("--limit", type=int, default=25,
                    help="rows to show in the text preview (default 25; 0 = all)")
    a = ap.parse_args(argv)

    hr = a.header_rows if a.header_rows == "auto" else int(a.header_rows)
    feed = parse_pos(a.path, sheet=a.sheet, header_rows=hr, tol=a.tol)

    if a.json:
        print(json.dumps({
            "path": feed.path, "sheet": feed.sheet, "header_rows": feed.header_rows,
            "ok": feed.ok,
            "records": [asdict(r) for r in feed.records],
            "mismatches": [asdict(m) for m in feed.mismatches],
            "warnings": feed.warnings,
        }, indent=2, default=str))
        return 0 if feed.ok else 1

    print(f"{feed.path}   sheet={feed.sheet!r}   header_rows={feed.header_rows}")
    print(f"  symbols parsed    : {len(feed.records)}")
    print(f"  col-15 mismatches : {len(feed.mismatches)}")
    print(f"  warnings          : {len(feed.warnings)}")
    for w in feed.warnings:
        print(f"    WARN  {w}")
    for m in feed.mismatches:
        d = "" if m.delta is None else f"  delta={m.delta:+,.2f}"
        note = f"  [{m.note}]" if m.note else ""
        print(f"    MISMATCH row {m.row:>4}  {m.symbol:<16} "
              f"expected={_fmt(m.expected)}  actual={_fmt(m.actual)}{d}{note}")

    rows = feed.records if a.limit == 0 else feed.records[:a.limit]
    print()
    for r in rows:
        flag = "" if r.verified else "  !UNVERIFIED"
        print(f"  {r.symbol:<16} LTP={_fmt(r.ltp)}  "
              f"net_series[{len(r.net_series)}]={r.net_series}{flag}")
    if a.limit and len(feed.records) > a.limit:
        print(f"  ... {len(feed.records) - a.limit} more (use --limit 0)")

    return 0 if feed.ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
